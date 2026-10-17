"""Every classification scheme in one format: schemes/<id>.json =
  {"id", "name", "source", "split": "train"|"heldout", "levels": [level names, top first],
   "nodes": [{"id", "name", "about"|None, "parent"|None, "level": 0..}]}

Real schemes from the web (raw/), our own map groups (search/fields/*/maps), invented schemes later (invent.py).
Held out entirely (never in training): Science-Metrix (the expert benchmark, jevarticle) and JEL (economics).

    nix shell --impure --expr 'with import (builtins.getFlake "nixpkgs") {}; python3.withPackages (p: [p.openpyxl])' \
        -c python3 build_schemes.py
"""
import csv, glob, html, json, re, xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

ROOT = Path("/mnt/fast/scholarsreadinglist/fieldjudge/v3/schemes")
RAW = ROOT / "raw"
HELDOUT = {"science-metrix", "jel"}


def save(sid, name, source, levels, nodes):
    seen = set()
    for n in nodes:
        assert n["id"] not in seen, (sid, n["id"]); seen.add(n["id"])
        n["name"] = re.sub(r"\s+", " ", html.unescape(n["name"])).strip()
        if n.get("about"):
            n["about"] = re.sub(r"\s+", " ", html.unescape(n["about"])).strip()[:300] or None
    ids = {n["id"] for n in nodes}
    assert all(n["parent"] is None or n["parent"] in ids for n in nodes), sid
    d = {"id": sid, "name": name, "source": source, "split": "heldout" if sid in HELDOUT else "train",
         "levels": levels, "nodes": nodes}
    (ROOT / f"{sid}.json").write_text(json.dumps(d, indent=0))
    per = defaultdict(int)
    for n in nodes:
        per[n["level"]] += 1
    print(f"{sid:28} {d['split']:8} " + " / ".join(f"{levels[k]} {per[k]}" for k in sorted(per)))


def node(i, name, parent, level, about=None):
    return {"id": str(i), "name": name, "about": about, "parent": None if parent is None else str(parent), "level": level}


# OpenAlex: domains > fields > subfields > topics (descriptions are OpenAlex's own, topics' generated)
L = lambda e: json.loads((RAW / f"openalex_{e}.json").read_text())
sid = lambda u: u["id"].rsplit("/", 1)[1]
nodes = [node("d" + sid(x), x["display_name"], None, 0, x.get("description")) for x in L("domains")]
nodes += [node("f" + sid(x), x["display_name"], "d" + sid(x["domain"]), 1, x.get("description")) for x in L("fields")]
nodes += [node("s" + sid(x), x["display_name"], "f" + sid(x["field"]), 2, x.get("description")) for x in L("subfields")]
nodes += [node("t" + sid(x), x["display_name"], "s" + sid(x["subfield"]), 3, x.get("description")) for x in L("topics")]
save("openalex", "OpenAlex topic classification", "https://openalex.org (CC0)", ["domain", "field", "subfield", "topic"], nodes)

# arXiv: groups > categories
s = (RAW / "arxiv.html").read_text()
nodes, groups = [], {}
for part in re.split(r'<h2 class="accordion-head"[^>]*>', s)[1:]:
    g = re.sub(r"<[^>]+>", " ", part.split("</h2>")[0]).strip()
    g = re.sub(r"\s+", " ", g)
    gid = "g" + str(len(groups)); groups[g] = gid
    nodes.append(node(gid, g, None, 0))
    for code, label, desc in re.findall(r"<h4>([\w.-]+) <span>\(([^)]*)\)</span></h4>\s*</div>\s*<div class=\"column\"><p>(.*?)</p>",
                                        part, re.S):
        nodes.append(node(code, label, gid, 1, re.sub(r"<[^>]+>", "", desc).split(". Roughly includes")[0]))
save("arxiv", "arXiv subject categories", "https://arxiv.org/category_taxonomy", ["group", "category"], nodes)

# ACM CCS 2012 (SKOS): three levels
ns = {"rdf": "http://www.w3.org/1999/02/22-rdf-syntax-ns#", "skos": "http://www.w3.org/2004/02/skos/core#"}
r = ET.parse(RAW / "acm.xml").getroot()
label, narrower = {}, defaultdict(list)
A = "{%s}about" % ns["rdf"]; R = "{%s}resource" % ns["rdf"]
tops = [t.get(R) for t in r.find("skos:ConceptScheme", ns).findall("skos:hasTopConcept", ns)]
for c in r.findall("skos:Concept", ns):
    label[c.get(A)] = c.find("skos:prefLabel", ns).text
    narrower[c.get(A)] = [n.get(R) for n in c.findall("skos:narrower", ns)]
nodes = []
def walk(c, parent, level):
    nodes.append(node(c, label[c], parent, level))
    if level < 2:
        for n in narrower[c]:
            walk(n, c, level + 1)
for t in tops:
    walk(t, None, 0)
save("acm-ccs", "ACM Computing Classification System 2012", "https://dl.acm.org/ccs", ["area", "topic", "subtopic"], nodes)

# JEL (held out)
r = ET.parse(RAW / "jel.xml").getroot()
nodes = []
def jwalk(e, parent):
    for c in e.findall("classification"):
        code, lvl = c.find("code").text, int(c.get("level")) - 1
        desc = re.sub(r"\s*&bull;\s*|\s*•\s*", "; ", html.unescape(c.find("description").text or ""))
        nodes.append(node(code, desc, parent, lvl))
        jwalk(c, code)
jwalk(r, None)
save("jel", "Journal of Economic Literature (JEL) classification", "https://www.aeaweb.org/econlit/jelCodes.php",
     ["category", "subcategory", "code"], nodes)

# MSC 2020: 2-digit areas > 3-character sections (top two levels; 5-character codes are too fine)
nodes = []
for row in csv.DictReader(open(RAW / "msc2020.csv", encoding="latin-1"), delimiter="\t"):
    code, text = row["code"].strip('"'), re.sub(r"\{\\(?:it|bf) ([^}]*)\}", r"\1", row["text"].strip('"'))
    text = re.sub(r"\[See also.*?\]|\{For .*?\}", "", text).strip(" ;")
    if re.fullmatch(r"\d\d-XX", code):
        nodes.append(node(code[:2], text, None, 0))
    elif re.fullmatch(r"\d\d[A-Z]xx", code) and code[:2] in {n["id"] for n in nodes}:
        nodes.append(node(code[:3], text, code[:2], 1))
save("msc2020", "Mathematics Subject Classification 2020", "https://msc2020.org (CC BY-NC-SA)", ["area", "section"], nodes)

# MeSH: 16 categories > top headings > their children (descriptor names; a heading in several trees appears under each)
CATS = {"A": "Anatomy", "B": "Organisms", "C": "Diseases", "D": "Chemicals and Drugs",
        "E": "Analytical, Diagnostic and Therapeutic Techniques, and Equipment", "F": "Psychiatry and Psychology",
        "G": "Phenomena and Processes", "H": "Disciplines and Occupations",
        "I": "Anthropology, Education, Sociology, and Social Phenomena", "J": "Technology, Industry, and Agriculture",
        "K": "Humanities", "L": "Information Science", "M": "Named Groups", "N": "Health Care",
        "V": "Publication Characteristics", "Z": "Geographicals"}
rows = json.loads((RAW / "mesh_depth2.json").read_text())
nodes = [node(k, v, None, 0) for k, v in CATS.items() if k != "V"]
done_tn = set()
for tn, name in sorted(rows, key=lambda x: (len(x[0]), x[0])):
    if tn[0] not in CATS or tn[0] == "V" or tn in done_tn:
        continue
    done_tn.add(tn)
    if False:
        continue
    if re.fullmatch(r"[A-Z]\d\d", tn):
        nodes.append(node(tn, name, tn[0], 1))
    elif re.fullmatch(r"[A-Z]\d\d\.\d{3}", tn) and tn[:3] in {n["id"] for n in nodes}:
        nodes.append(node(tn, name, tn[:3], 2))
save("mesh", "Medical Subject Headings (MeSH) tree", "https://www.nlm.nih.gov/mesh (public domain)",
     ["category", "heading", "subheading"], nodes)

# ANZSRC FoR 2020 (divisions > groups > fields, group definitions) and OECD FORD 2015 (from its correspondence table)
import openpyxl
wb = openpyxl.load_workbook(RAW / "anzsrc_for.xlsx", read_only=True)
nodes, have = [], set()
for r in wb["Table 3"].iter_rows(values_only=True):
    a, b, c, d = (list(r) + [None] * 4)[:4]
    if isinstance(a, int) and b:
        nodes.append(node(a, b.title() if b.isupper() else b, None, 0)); have.add(str(a))
    elif isinstance(b, int) and c and str(b)[:2] in have:
        nodes.append(node(b, c, str(b)[:2], 1)); have.add(str(b))
    elif isinstance(c, int) and d and str(c)[:4] in have:
        nodes.append(node(c, d, str(c)[:4], 2)); have.add(str(c))
defs = {}
for r in wb["Table 4"].iter_rows(values_only=True):
    cells = [x for x in r if x is not None]
    if len(cells) >= 3 and isinstance(cells[0], int) and len(str(cells[0])) == 4 and isinstance(cells[2], str):
        defs[str(cells[0])] = cells[2]
for n in nodes:
    n["about"] = (defs.get(n["id"]) or "").split(" This group excludes")[0][:300] or None
save("anzsrc-for", "ANZSRC Fields of Research 2020 (Australia and New Zealand)", "https://www.abs.gov.au (CC BY)",
     ["division", "group", "field"], nodes)
wb = openpyxl.load_workbook(RAW / "anzsrc_ford.xlsx", read_only=True)
nodes, top = [], None
for r in wb["Table 1"].iter_rows(values_only=True):
    a, b = (list(r) + [None, None])[:2]
    if isinstance(a, str) and re.match(r"^\d \w", a):
        top = a.split(" ", 1)[0]; nodes.append(node(top, a.split(" ", 1)[1].strip(), None, 0))
    if isinstance(b, str) and re.match(r"^\d\.\d+ ", b) and top:
        code, name = b.split(" ", 1)
        if code not in {n["id"] for n in nodes}:
            nodes.append(node(code, name.strip(), top, 1))
save("oecd-ford", "OECD Fields of Research and Development (FORD 2015, Frascati Manual)", "OECD, via ABS correspondence",
     ["field", "subfield"], nodes)

# Science-Metrix (held out: the expert benchmark)
tax = list(csv.DictReader(open(Path.home() / "Projects/jevarticle/data/taxonomy.csv")))
nodes, ids = [], {}
for t in tax:
    for lvl, key, parent in ((0, "domain", None), (1, "field", "domain"), (2, "subfield", "field")):
        k = (lvl, t[key])
        if k not in ids:
            ids[k] = f"{lvl}-{len(ids)}"
            nodes.append(node(ids[k], t[key], ids[(lvl - 1, t[parent])] if parent else None, lvl))
save("science-metrix", "Science-Metrix classification", "Rivest et al. 2021 (jevarticle)", ["domain", "field", "subfield"], nodes)

# Our own maps: a field's research groups (AI names and descriptions) > their subgroups. Newest version of each map.
n_maps = 0
for job in sorted(glob.glob("/mnt/fast/scholarsreadinglist/search/fields/*/")):
    vers = sorted(glob.glob(job + "maps/*/main.json"), key=lambda f: Path(f).stat().st_mtime)
    vers = [f for f in vers if all((g.get("label") or {}).get("name") for g in json.loads(Path(f).read_text())["groups"])]
    if not vers:
        continue
    m = json.loads(Path(vers[-1]).read_text())
    nodes = []
    for g in m["groups"]:
        nodes.append(node(f"g{g['id']}", g["label"]["name"], None, 0, g["label"].get("about")))
        sub = Path(vers[-1]).with_name(f"group-{g['id']}.json")
        if sub.exists():
            for h in json.loads(sub.read_text())["groups"]:
                if (h.get("label") or {}).get("name"):
                    nodes.append(node(f"g{g['id']}.{h['id']}", h["label"]["name"], f"g{g['id']}", 1, h["label"].get("about")))
    q = m.get("question") or Path(job).name
    save("map-" + Path(job).name[:12], f"Research groups of a map of the field: {q}", "our whole-field map (" + vers[-1].split("/")[-2] + ")",
         ["research group", "subgroup"], nodes)
    n_maps += 1
print("maps:", n_maps)
